from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[1]/'templates';root.mkdir(exist_ok=True)

def make():
    d=Document();s=d.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.top_margin=s.bottom_margin=Inches(.85);s.left_margin=s.right_margin=Inches(.9)
    for n in ['Normal','Title','Heading 1','Heading 2']:
        style=d.styles[n];style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0);style.font.size=Pt(12 if n=='Normal' else 18 if n=='Title' else 13)
        style.paragraph_format.space_after=Pt(7);style.paragraph_format.line_spacing=1.12
    for element in list(d.styles.element.iter()):
        if element.tag==qn('w:pBdr'):element.getparent().remove(element)
        if element.tag==qn('w:rFonts'):
            for key in list(element.attrib):
                if key.endswith('Theme'):del element.attrib[key]
    d.styles['Title'].paragraph_format.space_after=Pt(16)
    d.core_properties.author='RESMINAMA';d.core_properties.title='';d.core_properties.subject='';d.core_properties.comments=''
    return d

def p(d,text):return d.add_paragraph(text)

d=make();d.add_paragraph('Служебная записка','Title');p(d,'{{department}}');p(d,'От {{date}}    Номер {{number}}');p(d,'Кому: {{correspondent}}');p(d,'От: {{author}}');d.add_paragraph('{{title}}','Heading 1');p(d,'{{description}}');p(d,'Срок исполнения: {{due}}');p(d,'Автор: {{author}}');d.save(root/'memo.docx')
d=make();d.add_paragraph('Деловое письмо','Title');p(d,'{{department}}');p(d,'Получатель: {{correspondent}}');p(d,'Исходящий номер {{number}} от {{date}}');p(d,'На документ {{external_number}} от {{external_date}}');d.add_paragraph('{{title}}','Heading 1');p(d,'{{description}}');p(d,'Исполнитель: {{author}}');d.save(root/'letter.docx')
d=make();d.add_paragraph('Карточка документа','Title');d.add_paragraph('{{title}}','Heading 1');p(d,'Регистрационный номер: {{number}}');p(d,'Дата документа: {{date}}    Дата регистрации: {{registered_at}}');p(d,'Статус: {{status}}');p(d,'Корреспондент: {{correspondent}}');p(d,'Внешний номер: {{external_number}}    Дата: {{external_date}}');p(d,'Подразделение: {{department}}');p(d,'Автор: {{author}}');p(d,'Срок исполнения: {{due}}');p(d,'Категория: {{category}}');p(d,'Метки: {{tags}}');d.add_paragraph('Содержание','Heading 1');p(d,'{{description}}');d.add_paragraph('Финансовые реквизиты','Heading 1');p(d,'Сумма: {{amount}} {{currency}}');p(d,'Условия оплаты: {{payment}}');p(d,'Приоритет: {{priority}}');p(d,'Дело: {{case}}');d.add_page_break();d.add_paragraph('Работа с документом','Title');p(d,'{{number}} · {{title}}');d.add_paragraph('Участники и решения','Heading 1');p(d,'Редакторы: {{editors}}');p(d,'Согласование: {{approval_route}}');p(d,'Утверждающий: {{final_approver}}');d.add_paragraph('Вложения · последние версии','Heading 1');p(d,'{{files}}');d.add_paragraph('Поручения','Heading 1');p(d,'{{tasks}}');d.add_paragraph('Архив','Heading 1');p(d,'{{archive_info}}');d.save(root/'card.docx')
